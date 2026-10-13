"""Documents and code: PDF, Word, PowerPoint, text, Markdown, code.

Each file is split into ~1200-character passages (150 overlap, paragraph-aware) and each
passage is embedded as "title: <file name> | text: <passage>". A file's score is its best
passage, and that passage is shown as the snippet.
"""
from __future__ import annotations

import html
import os
import re
import zipfile

import numpy as np

from .. import config, secretmask
from ..embedder import doc
from .base import Result, Source, open_pdf, pretty_dir
from . import walk as walk_mod
from .walk import ext_of, sync

TEXT = {".env", ".txt", ".md", ".markdown", ".rst", ".org", ".tex", ".csv", ".log", ".json", ".yaml", ".yml", ".toml",
        ".ini", ".html", ".htm", ".xml"}
CODE = {".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp", ".cc", ".h", ".hpp", ".cs", ".go", ".rs",
        ".rb", ".php", ".swift", ".kt", ".sql", ".sh", ".ps1", ".bat", ".css", ".scss", ".vue", ".ipynb", ".r", ".m"}
DOCS = {".pdf", ".docx", ".pptx"}
EXTS = TEXT | CODE | DOCS
# Machine-written data: the first chunks say what the file is; embedding all of it costs hours and finds nothing.
DATA = {".json", ".csv", ".log", ".xml", ".yaml", ".yml", ".toml", ".ini"}
DENSE_CANDIDATES = 300  # embedding ranking: this many best chunks join the keyword hits
RRF_K = 60              # reciprocal rank fusion constant (the usual value; rank 0 vs 10 matters, 200 vs 210 doesn't)
MAX_BYTES = 3_000_000  # bigger text files (dataset dumps) are indexed by their head only
HEAD_BYTES = 64_000
LANG = {".py": "Python", ".js": "JavaScript", ".ts": "TypeScript", ".tsx": "TSX", ".jsx": "JSX", ".java": "Java",
        ".c": "C", ".cpp": "C++", ".cc": "C++", ".h": "C", ".hpp": "C++", ".cs": "C#", ".go": "Go", ".rs": "Rust",
        ".rb": "Ruby", ".php": "PHP", ".swift": "Swift", ".kt": "Kotlin", ".sql": "SQL", ".sh": "Shell",
        ".ps1": "PowerShell", ".bat": "Batch", ".css": "CSS", ".scss": "SCSS", ".vue": "Vue", ".r": "R", ".m": "MATLAB"}
ICON = {".pdf": "📕", ".docx": "📘", ".pptx": "📙", ".md": "📝", ".txt": "📄", ".env": "🔑"}


# ---- text extraction -------------------------------------------------------------------

def extract(path: str) -> list[tuple[str, str]]:
    """[(location label, text)], e.g. [("p. 3", "...")] for PDFs, [("", whole text)] for plain files."""
    ext = ext_of(path)
    if ext == ".pdf":
        from pypdf import PdfReader
        r = PdfReader(path)
        return [(f"p. {i + 1}", pg.extract_text() or "") for i, pg in enumerate(r.pages[:300])]
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            parts += [" | ".join(c.text for c in row.cells) for row in t.rows]
        return [("", "\n".join(parts))]
    if ext == ".pptx":
        out = []
        with zipfile.ZipFile(path) as z:
            slides = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                            key=lambda n: int(re.findall(r"\d+", n)[-1]))
            for i, n in enumerate(slides):
                xml = z.read(n).decode("utf-8", "ignore")
                out.append((f"slide {i + 1}", " ".join(html.unescape(t) for t in re.findall(r"<a:t>([^<]*)</a:t>", xml))))
        return out
    size = os.path.getsize(path)
    if ext == ".ipynb" and size > 10 * MAX_BYTES:
        return []  # notebooks are big because of embedded outputs, but past ~30 MB it's not worth parsing
    head = ext in DATA or (size > MAX_BYTES and ext != ".ipynb")
    with open(path, "rb") as f:
        raw = f.read(HEAD_BYTES if head else -1)
    if b"\0" in raw[:4096]:
        return []  # binary file with a text extension
    text = raw.decode("utf-8", "ignore")
    if ext == ".ipynb":
        import json
        try:
            cells = json.loads(text).get("cells", [])
            text = "\n\n".join("".join(c.get("source", "")) for c in cells)
        except ValueError:
            return []
    elif ext in (".html", ".htm", ".xml"):
        text = html.unescape(re.sub(r"(?s)<(script|style).*?</\1>|<[^>]+>", " ", text))
    return [("", text)]


def chunks(parts: list[tuple[str, str]], size: int, overlap: int, limit: int) -> list[tuple[str, str]]:
    """Paragraph-aware windows of ~size chars. Keeps the location label of where each chunk starts."""
    out: list[tuple[str, str]] = []
    for loc, text in parts:
        text = re.sub(r"[ \t]+", " ", text).strip()
        if len(text) < 20:
            continue
        paras = re.split(r"\n\s*\n|(?<=[.!?])\s+(?=[A-Z])", text)
        buf = ""
        for p in paras:
            if len(buf) + len(p) > size and buf:
                out.append((loc, buf.strip()))
                buf = buf[-overlap:] + " "
            while len(p) > size:  # very long paragraph / minified line
                out.append((loc, (buf + p[: size - len(buf)]).strip()))
                p = p[size - len(buf) - overlap:]
                buf = ""
            buf += p + "\n"
        if buf.strip():
            out.append((loc, buf.strip()))
        if len(out) >= limit:
            break
    return [(l, t) for l, t in out if len(t) > 30][:limit]


class FilesSource(Source):
    name = "files"
    label = "Documents & code"
    panel = "list"
    exts, min_size = EXTS, 20
    intent = ("A document, PDF, note, report, slide deck or source code file on my computer, found by what it "
              "says or is about, including config and .env files where I saved an API key, token or password.")
    supports_image_query = True
    floor, strong = 0.675, 0.80   # noise band on test corpus: 0.61-0.67
    image_floor, image_strong = 0.72, 0.80  # image->text noise reached 0.69
    # Real index (49.8k chunks): nonsense queries' best hit p50 0.748, p90 0.772; real queries 0.789-0.834.
    floor_ref_n, floor_per_decade = 100, 0.039

    def index(self, emb, progress) -> None:
        sync(self.store, self.name, self.exts, self.handler(emb), progress, min_size=self.min_size, label="documents")
        if walk_mod.RECENT_SINCE is None:
            self._fulltext_pass(progress)

    def handler(self, emb):
        """handle(path) -> status: index one file (its old items are already removed)."""
        size, overlap = config.get("chunk_chars"), config.get("chunk_overlap")

        def handle(path: str) -> str:
            cs = chunks(extract(path), size, overlap, _limit(path))
            if not cs:
                return "empty"
            name = os.path.basename(path)
            vecs = emb.texts([doc(t, name) for _, t in cs])
            for i, ((loc, t), v) in enumerate(zip(cs, vecs)):
                self.store.add(self.name, f"files:{path}#{i}", v, path=path, title=name,
                               subtitle=os.path.dirname(path), snippet=t, meta={"loc": loc}, commit=False)
            return "ok"

        return handle

    def _fulltext_pass(self, progress) -> None:
        """Once: documents indexed before the keyword index kept only their first 600 characters per chunk.
        Re-read them (no GPU) and store each chunk's full text, so BM25 sees every word."""
        if self.store.kv_get("fulltext") == "done":
            return
        size, overlap = config.get("chunk_chars"), config.get("chunk_overlap")
        paths = sorted(p for p, st in self.store.file_status(self.name).items() if st == "ok")
        start = int(self.store.kv_get("fulltext_at") or 0)
        for i in range(start, len(paths)):
            walk_mod.gate(progress)
            if i % 25 == 0:
                progress(f"{i:,}/{len(paths):,} documents: full text for keyword search", i, len(paths))
                self.store.kv_set("fulltext_at", str(i))
            p = paths[i]
            try:
                cs = chunks(extract(p), size, overlap, _limit(p))
            except Exception:  # unreadable now: keeps its 600-character text
                continue
            self.store.set_snippets([(t, f"files:{p}#{j}") for j, (_, t) in enumerate(cs)])
        self.store.kv_set("fulltext", "done")

    def search(self, emb, qvec, text, k=8) -> list[Result]:
        """Hybrid: the embedding ranking and the BM25 keyword ranking fused (reciprocal rank fusion), so a
        document is found by what it means and by the exact words in it. Each result carries `lexical`
        (share of the query's information it contains) for the section's keep rule in core."""
        ids, s = self.store.scores(self.name, qvec)  # scope / date window / file type already applied
        if not len(ids):
            return []
        order = np.argsort(-s)
        score = {int(ids[i]): float(s[i]) for i in order[:DENSE_CANDIDATES]}
        rank = {i: r for r, i in enumerate(score)}
        lex = self.store.lexical(text) if text else {}
        weights = self.store.term_idf(text) if text else {}  # which query words to centre snippets on
        missing = [i for i in lex if i not in score]
        if missing:  # keyword hits the embedding ranked low: their dense score and rank
            desc = -s[order]
            for p in np.nonzero(np.isin(ids, np.array(missing, np.int64)))[0]:
                i = int(ids[p])
                score[i], rank[i] = float(s[p]), int(np.searchsorted(desc, -s[p]))
        fused = {i: 1.0 / (RRF_K + rank[i]) + (1.0 / (RRF_K + lex[i][0]) if i in lex else 0.0)
                 for i in score if score[i] > -1.0}  # -2 marks items outside the scope / window / type
        out, seen = [], set()
        for item_id in sorted(fused, key=lambda i: -fused[i]):
            it = self.store.get(item_id)
            if not it or it.path in seen:
                continue
            seen.add(it.path)
            ext = ext_of(it.path)
            loc = it.meta.get("loc", "")
            kind = _kind(ext)
            meta = {"ext": ext}
            if ext in LANG:
                meta["lang"] = LANG[ext]
            page = re.fullmatch(r"p\. (\d+)", loc)
            if page:
                meta["page"] = int(page.group(1))
            # Keys and passwords show as AIza••••••3kQ; "Copy key" copies the first one in the snippet.
            passage = secretmask.mask(it.snippet)
            snippet = _code_snip(passage, weights) if kind == "code" else _snip(passage, weights)
            actions = [{"label": "Show in folder", "op": "reveal"}]
            if passage != it.snippet:
                shown = [i for i, (a, b) in enumerate(secretmask.find(it.snippet))
                         if secretmask.hide(it.snippet[a:b]) in snippet]
                if shown:
                    actions.insert(0, {"label": "Copy key", "op": f"copy_secret:{shown[0]}"})
                    meta["secret"] = True
            out.append(Result(
                id=f"files:{it.id}", source=self.name, title=it.title, subtitle=pretty_dir(it.subtitle),
                snippet=snippet,
                score=score[item_id], icon=ICON.get(ext, "🧾" if ext in CODE else "📄"),
                badge=(ext[1:].upper() + (f" · {loc}" if loc else "")), kind=kind, meta=meta,
                # PDF pages are rendered on demand by /api/preview (no cost at index time).
                thumb=f"/api/preview/{it.id}" if page else None, actions=actions))
            # server-side only (not dataclass fields, so never sent to the UI)
            out[-1].passage = it.snippet
            out[-1].fused = fused[item_id]
            out[-1].lexical, out[-1].lexical_info = (lex[item_id][1], lex[item_id][2]) if item_id in lex else (0.0, 0.0)
            if len(out) >= k:
                break
        return out

    def run(self, result_id: str, op: str = "open") -> dict:
        it = self.store.get(int(result_id.split(":", 1)[1]))
        if it is not None and op.startswith("copy_secret:"):
            vals = secretmask.values(it.snippet)
            n = int(op.split(":", 1)[1])
            if n >= len(vals):
                return {"ok": False, "error": "key not found"}
            return {"ok": True, "copy": vals[n], "toast": "Key copied"}
        if it is not None and op == "open" and it.path and it.path.lower().endswith(".pdf"):
            m = re.fullmatch(r"p\. (\d+)", it.meta.get("loc", ""))
            open_pdf(it.path, int(m.group(1)) if m else None)  # at the matching page
            return {"ok": True}
        return super().run(result_id, op)


def _limit(path: str) -> int:
    """Chunks kept per file: machine data and huge text files are found by their head."""
    ext = ext_of(path)
    if ext in DATA or os.path.getsize(path) > MAX_BYTES:
        return config.get("max_chunks_data")
    if ext in CODE or ext in (".html", ".htm"):
        return config.get("max_chunks_code")
    return config.get("max_chunks_per_file")


def _kind(ext: str) -> str:
    return {".pdf": "pdf", ".docx": "doc", ".pptx": "slides", ".ipynb": "notebook"}.get(ext) or (
        "code" if ext in CODE else "data" if ext in DATA or ext == ".env" else "text")


def _word_re(term: str) -> re.Pattern:
    """A query word at the start of a word in the text; "awards" also finds "award" (the index stems too)."""
    stem = term[:-1] if len(term) > 3 and term.endswith("s") else term
    return re.compile(r"(?<!\w)" + re.escape(stem), re.I)


def _code_snip(passage: str, weights: dict[str, float], lines: int = 4) -> str:
    """A few whole lines of code, starting just above the line with the most query information."""
    ls = [l.rstrip() for l in passage.splitlines() if l.strip()]
    pats = [(_word_re(w), v) for w, v in weights.items()]
    gain = [sum(v for p, v in pats if p.search(l)) for l in ls]
    start = max(range(len(ls)), key=lambda i: (gain[i], -i)) if ls and any(gain) else 0
    return "\n".join(l[:120] for l in ls[max(0, start - 1):max(0, start - 1) + lines])


def _snip(passage: str, weights: dict[str, float], width: int = 240) -> str:
    """The `width` characters of the passage holding the most query information: each distinct query word
    found counts its IDF, so a window with "contoso" and "award" beats one with "the" five times."""
    p = re.sub(r"\s+", " ", passage).strip()
    if len(p) <= width:
        return p
    hits = sorted((m.start(), w) for w in weights for m in _word_re(w).finditer(p))
    best, best_key = 0, (0.0, 0.0)
    for a, anchor in hits:  # windows that open a little before a hit; ties go to the rarest word up front
        start = max(0, a - width // 4)
        gain = sum(weights[w] for w in {w for b, w in hits if start <= b < start + width - 10})
        key = (round(gain, 6), weights[anchor])
        if key > best_key:
            best, best_key = start, key
    if best:  # begin on a word boundary
        sp = p.find(" ", best)
        best = sp + 1 if 0 <= sp < best + 20 else best
    return ("…" if best else "") + p[best:best + width].rstrip() + ("…" if best + width < len(p) else "")
